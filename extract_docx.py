import docx, sys, glob, os
sys.stdout = open(r"C:\Users\itren\Desktop\docx_output.txt", "w", encoding="utf-8")

files = [
    r"C:\Users\itren\Desktop\موقع بحر الشهر\من نحن.docx",
    r"C:\Users\itren\Desktop\موقع بحر الشهر\المقاولات.docx",
    r"C:\Users\itren\Desktop\موقع بحر الشهر\العقار.docx",
    r"C:\Users\itren\Desktop\موقع بحر الشهر\محتوى موقع بحر الشهر.docx",
    r"C:\Users\itren\Desktop\موقع بحر الشهر\بعد التعديلات\من نحن.docx",
    r"C:\Users\itren\Desktop\موقع بحر الشهر\بعد التعديلات\المقاولات (1).docx",
    r"C:\Users\itren\Desktop\موقع بحر الشهر\الصفحة الرئيسية.docx",
]

for f in files:
    print("="*80)
    print("FILE:", f)
    print("="*80)
    try:
        d = docx.Document(f)
        for p in d.paragraphs:
            t = p.text.strip()
            if t:
                print(t)
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                print(" | ".join(cells))
    except Exception as e:
        print("ERROR:", e)
    print()
